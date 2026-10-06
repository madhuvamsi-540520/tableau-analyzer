"""Export a comprehensive multi-sheet .xlsx inventory (openpyxl)."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1F6FEB")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb: Workbook, title: str, headers: list, rows: list) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)] + [0])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 60)


def build_xlsx(metadata: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    counts = metadata.get("counts", {})
    cx = metadata.get("complexity", {})
    summary = metadata.get("assessment", {}).get("summary", {})
    _sheet(wb, "Summary", ["Metric", "Value"], [
        ["Workbook", metadata.get("name")],
        ["Version", metadata.get("version")],
        ["Source Platform", metadata.get("sourcePlatform")],
        ["Data Sources", counts.get("dataSources")],
        ["Tables", counts.get("tables")],
        ["Columns", counts.get("columns")],
        ["Relationships", counts.get("relationships")],
        ["Joins", cx.get("joins")],
        ["Calculated Fields", cx.get("calculatedFields")],
        ["Filters", cx.get("filters")],
        ["Parameters", cx.get("parameters")],
        ["Custom SQL", cx.get("customSql")],
        ["Maturity Score", summary.get("maturityScore")],
        ["Maturity Band", summary.get("maturityBand")],
    ])

    conn, tbl, col, rel, flt, sql, cls = [], [], [], [], [], [], []
    for ds in metadata.get("dataSources", []):
        dsn = ds.get("caption") or ds.get("name")
        c = ds.get("classification", {})
        cls.append([dsn, c.get("type"), ", ".join(c.get("factTables", [])),
                    ", ".join(c.get("dimensionTables", [])), " ".join(c.get("reasoning", []))])
        for cn in ds.get("connections", []):
            conn.append([dsn, cn.get("friendlyType"), cn.get("class"), cn.get("server") or "",
                         cn.get("database") or "", cn.get("schema") or "", cn.get("warehouse") or "",
                         cn.get("fileLocation") or ""])
        for t in ds.get("tables", []):
            tbl.append([dsn, t.get("name"), t.get("kind"), t.get("schema") or "",
                        t.get("database") or "", t.get("columnCount"),
                        "" if t.get("rowCount") is None else t.get("rowCount")])
            for cc in t.get("columns", []):
                col.append([dsn, t.get("name"), cc.get("name"), cc.get("originalName"),
                            cc.get("dataType"), cc.get("role"), cc.get("defaultAggregation") or "",
                            cc.get("hidden"), cc.get("nullable")])
        for r in ds.get("relationships", []):
            rel.append([dsn, r.get("kind"), r.get("parentTable"), r.get("childTable"),
                        r.get("joinType") or "", r.get("cardinality") or "",
                        "; ".join(r.get("joinKeys", []))])
        for f in ds.get("filters", []):
            flt.append([dsn, f.get("field"), f.get("filterType"), f.get("filterClass"),
                        "; ".join(f.get("values", [])), f.get("expression")])
        for s in ds.get("customSql", []):
            a = s.get("analysis", {})
            sql.append([dsn, s.get("name"), s.get("sql"), ", ".join(a.get("referencedTables", [])),
                        ", ".join(a.get("aggregations", [])), a.get("hasUnionAll"), a.get("subqueryCount")])

    _sheet(wb, "Classification", ["Data Source", "Model", "Fact Tables", "Dimension Tables", "Reasoning"], cls)
    _sheet(wb, "Connections", ["Data Source", "Type", "Class", "Server", "Database", "Schema", "Warehouse", "File"], conn)
    _sheet(wb, "Tables", ["Data Source", "Table", "Kind", "Schema", "Database", "Columns", "Rows"], tbl)
    _sheet(wb, "Columns", ["Data Source", "Table", "Column", "Original", "Data Type", "Role", "Default Agg", "Hidden", "Nullable"], col)
    if rel:
        _sheet(wb, "Relationships", ["Data Source", "Kind", "Parent", "Child", "Join Type", "Cardinality", "Keys"], rel)
    if flt:
        _sheet(wb, "Filters", ["Data Source", "Field", "Type", "Class", "Values", "Expression"], flt)
    if sql:
        _sheet(wb, "Custom SQL", ["Data Source", "Name", "SQL", "Referenced Tables", "Aggregations", "UNION ALL", "Subqueries"], sql)

    assessment = metadata.get("assessment", {})
    findings = []
    for key, label in [
        ("technicalDebt", "Technical Debt"), ("modelingRisks", "Modeling Risk"),
        ("dataQualityRisks", "Data Quality"), ("governance", "Governance"),
        ("scalability", "Scalability"), ("performance", "Performance"),
        ("optimization", "Optimization"),
    ]:
        for f in assessment.get(key, []):
            findings.append([label, f.get("severity"), f.get("title"), f.get("businessImpact"),
                             f.get("recommendation"), "; ".join(f.get("evidence", []))])
    if findings:
        _sheet(wb, "Assessment", ["Section", "Severity", "Finding", "Business Impact", "Recommendation", "Evidence"], findings)

    frameworks = assessment.get("bestPractices", {}).get("frameworks", [])
    if frameworks:
        _sheet(wb, "Best Practices", ["Framework", "Score", "Notes"],
               [[f.get("name"), f.get("score"), f.get("notes")] for f in frameworks])

    out = BytesIO()
    wb.save(out)
    return out.getvalue()

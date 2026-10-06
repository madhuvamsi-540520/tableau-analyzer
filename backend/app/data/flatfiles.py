"""Best-effort row counts from flat files bundled inside a .twbx package.

Row counts are informational ("if available"): CSV/TXT are counted by lines,
XLSX/XLSM by streaming rows. Anything unreadable is skipped silently — a missing
count is never an error. Keys are the file stem / sheet name, lowercased, so the
parser can match them to table names.
"""
from __future__ import annotations

import csv  # noqa: F401  (kept for parity / future dialect handling)
import io
import zipfile
from pathlib import PurePosixPath


def row_counts_from_twbx(data: bytes) -> dict[str, int]:
    counts: dict[str, int] = {}
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return counts

    for info in zf.infolist():
        if info.filename.endswith("/"):
            continue
        path = PurePosixPath(info.filename)
        ext = path.suffix.lower()
        stem = path.stem.lower()
        try:
            if ext in (".csv", ".txt"):
                counts[stem] = _count_csv(zf, info)
            elif ext in (".xlsx", ".xlsm"):
                _count_xlsx(zf, info, counts)
        except Exception:
            # Best-effort only: a file we cannot read simply has no row count.
            continue
    return counts


def _count_csv(zf: zipfile.ZipFile, info: zipfile.ZipInfo) -> int:
    with zf.open(info) as fh:
        wrapper = io.TextIOWrapper(fh, encoding="utf-8", errors="replace", newline="")
        n = sum(1 for _ in wrapper)
    return max(n - 1, 0)  # subtract the header row


def _count_xlsx(zf: zipfile.ZipFile, info: zipfile.ZipInfo, counts: dict[str, int]) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return
    with zf.open(info) as fh:
        blob = fh.read()
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = sum(1 for _ in ws.iter_rows(values_only=True))
            counts[ws.title.lower()] = max(rows - 1, 0)
    finally:
        wb.close()

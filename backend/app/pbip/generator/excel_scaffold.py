"""Sample data + companion Excel template for tables with no readable data.

Ported from obiee-analyzer's ``generator/excel_scaffold.py``: every table must
be embedded inline (``Table.FromRows``) so the PBIP opens with no external file
or path. When the workbook's real data cannot be read (e.g. it only exists in a
``.hyper`` extract, or a .twb points at a file on another machine), the table
gets deterministic, typed sample rows. Relationship keys get consistent values
across related tables so joined visuals actually return rows.

The companion workbook carries the same schema (one sheet per table, exact
source-column names) so the user can paste real data in later.
"""
from __future__ import annotations

import datetime as _dt
import io

from openpyxl import Workbook

from ..model.ir import Column, Model, Table

DIM_ROWS = 8         # rows per dimension (one-side) table
FACT_ROWS = 24       # rows per fact (many-side) table
_BASE_DATE = _dt.datetime(2024, 1, 1)


def _key_sets(model: Model) -> tuple[set[tuple[str, str]], dict[tuple[str, str], str]]:
    """(primary-key set, foreign-key -> referenced-table map) keyed by (table, column)."""
    pk: set[tuple[str, str]] = set()
    fk: dict[tuple[str, str], str] = {}
    for r in model.relationships:
        pk.add((r.to_table, r.to_column))
        fk[(r.from_table, r.from_column)] = r.to_table
    return pk, fk


def _row_count(table: Table, fk: dict[tuple[str, str], str]) -> int:
    is_many_side = any(t == table.name for (t, _c) in fk)
    return FACT_ROWS if is_many_side else DIM_ROWS


def _sample(col: Column, i: int, table_name: str, pk: set, fk: dict) -> object:
    tc = (table_name, col.name)
    if tc in pk:
        return i + 1                                  # unique key 1..N
    if tc in fk:
        return (i % DIM_ROWS) + 1                     # FK cycles over the dim's keys
    dt = col.data_type
    if dt == "int64":
        return i + 1
    if dt == "double":
        return round(100.0 + i * 12.5, 2)
    if dt == "dateTime":
        return _BASE_DATE + _dt.timedelta(days=i * 30)
    if dt == "boolean":
        return i % 2 == 0
    return f"{col.name} {i + 1}"


def sample_rows(model: Model) -> dict[str, tuple[list[str], list[list]]]:
    """Deterministic sample data per table: {table_name: (source_columns, rows)}."""
    pk, fk = _key_sets(model)
    out: dict[str, tuple[list[str], list[list]]] = {}
    for table in model.tables:
        cols = [c.source_column for c in table.columns]
        n = _row_count(table, fk)
        rows = [[_sample(c, i, table.name, pk, fk) for c in table.columns] for i in range(n)]
        out[table.name] = (cols, rows)
    return out


def build_workbook(model: Model) -> bytes:
    """.xlsx bytes: one sheet per table with exactly the data embedded in the model
    (real where it was available, sample otherwise)."""
    wb = Workbook()
    wb.remove(wb.active)

    used: set[str] = set()
    for table in model.tables:
        sheet = (table.source_name or table.name)[:31] or "Sheet"
        base, n = sheet, 1
        while sheet in used:                          # openpyxl forbids duplicate sheet names
            suffix = f"_{n}"
            sheet = base[:31 - len(suffix)] + suffix
            n += 1
        used.add(sheet)
        ws = wb.create_sheet(title=sheet)

        cols = table.inline_columns or [c.source_column for c in table.columns]
        if not cols:
            continue
        ws.append(cols)
        for row in table.inline_rows or []:
            ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

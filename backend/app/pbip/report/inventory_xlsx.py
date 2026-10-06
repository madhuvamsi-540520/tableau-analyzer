"""Build an Excel (.xlsx) inventory report from the migration IR.

Sheets: Summary, Tables, Relationships, Calculated Fields. Human-readable audit
of what the workbook contains and how each calc field maps to DAX.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..model.ir import Model

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="0F1A3C")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _cardinality(from_cardinality: str) -> str:
    return "One-to-one" if from_cardinality == "one" else "Many-to-one"


def _crossfilter(cross_filter: str) -> str:
    return "Both" if cross_filter == "bothDirections" else "Single"


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    ws.freeze_panes = "A2"


def _autosize(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_inventory_xlsx(model: Model) -> bytes:
    key_cols = set()
    for r in model.relationships:
        key_cols.add((r.from_table, r.from_column))
        key_cols.add((r.to_table, r.to_column))

    wb = Workbook()

    # ---- Summary ----
    ws = wb.active
    ws.title = "Summary"
    conn = next(iter(model.connections.values()), None)
    source = ""
    if conn:
        source = conn.filename or f"{conn.cls} {conn.server or ''} {conn.database or ''}".strip()
    rows = [
        ("Project", model.name),
        ("Source caption", model.source_caption or ""),
        ("Source connection", source),
        ("Tables", len(model.tables)),
        ("Columns", sum(len(t.columns) for t in model.tables)),
        ("Relationships", len(model.relationships)),
        ("Calculated fields", len(model.calculated_fields)),
        ("  translated to DAX", sum(1 for cf in model.calculated_fields if cf.translated)),
        ("Report pages (tables)", len(model.worksheets)),
    ]
    _write_header(ws, ["Item", "Value"])
    for label, value in rows:
        ws.append([label, value])
    _autosize(ws, [26, 70])

    # ---- Tables ----
    ws = wb.create_sheet("Tables")
    _write_header(ws, ["Table", "Kind", "Column", "Data Type", "Summarize By", "Key"])
    for t in model.tables:
        if not t.columns:
            ws.append([t.name, t.kind, "", "", "", ""])
        for c in t.columns:
            ws.append([
                t.name, t.kind, c.name, c.data_type, c.summarize_by,
                "Yes" if (t.name, c.name) in key_cols else "",
            ])
    _autosize(ws, [22, 14, 26, 12, 14, 6])

    # ---- Relationships ----
    ws = wb.create_sheet("Relationships")
    _write_header(ws, ["From Table", "From Key", "", "To Table", "To Key", "Type", "Cross-filter"])
    for r in model.relationships:
        ws.append([
            r.from_table, r.from_column, "->", r.to_table, r.to_column,
            _cardinality(r.from_cardinality), _crossfilter(r.cross_filter),
        ])
    if not model.relationships:
        ws.append(["(none)", "", "", "", "", "", ""])
    _autosize(ws, [22, 20, 5, 22, 20, 14, 12])

    # ---- Calculated Fields ----
    ws = wb.create_sheet("Calculated Fields")
    _write_header(ws, ["Name", "Tableau Formula", "Equivalent DAX", "Method", "Review?"])
    for cf in model.calculated_fields:
        ws.append([
            cf.name,
            cf.formula,
            cf.dax or "(not translated - BLANK())",
            {"rule": "Offline rules", "llm": "LLM"}.get(cf.method, "-"),
            "Yes" if (not cf.translated or cf.needs_review) else "",
        ])
    if not model.calculated_fields:
        ws.append(["(none)", "", "", "", ""])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):
        for cell in row:
            cell.alignment = _WRAP
    _autosize(ws, [24, 46, 52, 14, 8])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
